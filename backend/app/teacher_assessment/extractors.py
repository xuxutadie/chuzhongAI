"""本地受限文档适配器；仅返回候选素材，不识别答案、不联网、不写文件。"""
from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
from time import monotonic
import math
import warnings
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from PIL import Image, ImageOps
from pypdf import PdfReader
import pypdfium2 as pdfium

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
IMAGE_FORMATS = {'image/jpeg': 'JPEG', 'image/png': 'PNG', 'image/webp': 'WEBP'}
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


class DocumentError(ValueError):
    """可向教师显示的错误；不得携带内部路径或第三方异常详情。"""


@dataclass(frozen=True)
class DocumentLimits:
    max_document_bytes: int = 20 * 1024 * 1024
    max_image_bytes: int = 5 * 1024 * 1024
    max_pdf_pages: int = 50
    max_docx_paragraphs: int = 20000
    max_zip_entries: int = 2000
    max_uncompressed_bytes: int = 100 * 1024 * 1024
    max_image_pixels: int = 20_000_000
    max_render_pixels: int = 4_000_000
    max_output_bytes: int = 100 * 1024 * 1024
    max_text_characters: int = 1_000_000
    max_seconds: float = 60


@dataclass(frozen=True)
class ExtractedAsset:
    data: bytes
    media_type: str
    width: int
    height: int


@dataclass
class ExtractedSource:
    index: int
    text: str
    formulas: list[str] = field(default_factory=list)
    asset_refs: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class ExtractedDocument:
    index_kind: str
    sources: list[ExtractedSource] = field(default_factory=list)
    assets: dict[str, ExtractedAsset] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


class _Budget:
    def __init__(self, limits):
        self.limits = limits
        self.started = monotonic()
        self.output_bytes = 0
        self.text_characters = 0

    def check(self):
        if monotonic() - self.started > self.limits.max_seconds:
            raise DocumentError('解析超时，请拆分文件后重试。')

    def text(self, value):
        self.check()
        self.text_characters += len(value)
        if self.text_characters > self.limits.max_text_characters:
            raise DocumentError('文字内容过多，请拆分文件后重试。')
        return value

    def asset(self, result, image):
        self.check()
        # 重建像素图，而非保存原对象；不带入EXIF、ICC或原文件附加内容。
        image = ImageOps.exif_transpose(image)
        clean = Image.new('RGB', image.size, 'white')
        rgba = image.convert('RGBA')
        clean.paste(rgba, mask=rgba.getchannel('A'))
        target = BytesIO()
        clean.save(target, format='PNG')
        content = target.getvalue()
        ref = sha256(content).hexdigest()
        if ref not in result.assets:
            self.output_bytes += len(content)
            if self.output_bytes > self.limits.max_output_bytes:
                raise DocumentError('解析后的图片过大，请拆分文件。')
            result.assets[ref] = ExtractedAsset(content, 'image/png', clean.width, clean.height)
        return ref


def _image_asset(data, result, budget, expected_format=None):
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as image:
            if image.format not in IMAGE_FORMATS.values() or (expected_format and image.format != expected_format):
                raise DocumentError('图片内容与文件类型不符，请上传 JPEG、PNG 或 WebP。')
            if image.width * image.height > budget.limits.max_image_pixels:
                raise DocumentError('图片像素过大，请缩小后上传。')
            if getattr(image, 'n_frames', 1) != 1:
                raise DocumentError('暂不支持动画图片，请上传单张静态图片。')
            image.verify()
        with Image.open(BytesIO(data)) as image:
            return budget.asset(result, image)


def _pdf(data, budget):
    if not data.startswith(b'%PDF-'):
        raise DocumentError('PDF 文件签名不匹配。')
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise DocumentError('PDF 带有密码，请解除密码后重新上传。')
    count = len(reader.pages)
    if not 1 <= count <= budget.limits.max_pdf_pages:
        raise DocumentError('PDF 页数超出限制，请分批上传。')
    root = reader.trailer['/Root']
    names = root.get('/Names', {})
    if hasattr(names, 'get_object'):
        names = names.get_object()
    if '/JavaScript' in names or '/OpenAction' in root or '/AA' in root:
        raise DocumentError('PDF 含自动动作，请导出为普通 PDF 后上传。')
    result = ExtractedDocument(index_kind='page')
    with pdfium.PdfDocument(data) as document:
        if len(document) != count:
            raise DocumentError('PDF 页结构异常，无法可靠提取。')
        for index, source in enumerate(reader.pages, 1):
            budget.check()
            text = budget.text(source.extract_text() or '')
            page = document[index-1]
            try:
                width, height = page.get_size()
                if not math.isfinite(width * height) or min(width, height) <= 0:
                    raise DocumentError('PDF 页面尺寸异常。')
                scale = min(2, math.sqrt(budget.limits.max_render_pixels / (width * height)))
                bitmap = page.render(scale=scale)
                try:
                    ref = budget.asset(result, bitmap.to_pil())
                finally:
                    bitmap.close()
            finally:
                page.close()
            # 即使有文字层也保留渲染图，避免几何标注/公式丢失或错序。
            issues = ['verify_visual_math']
            if not text.strip():
                issues.append('needs_ocr')
            result.sources.append(ExtractedSource(index, text, asset_refs=[ref], warnings=issues))
    return result


def _checked_docx(data, budget):
    if not data.startswith(b'PK\x03\x04'):
        raise DocumentError('Word 文件不是有效 DOCX，请勿仅修改扩展名。')
    with zipfile.ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > budget.limits.max_zip_entries or sum(item.file_size for item in entries) > budget.limits.max_uncompressed_bytes:
            raise DocumentError('Word 解压规模超出限制，请拆分文件。')
        seen = set()
        for item in entries:
            budget.check()
            path = PurePosixPath(item.filename)
            if path.is_absolute() or '..' in path.parts or '\\' in item.filename or ':' in item.filename or item.filename in seen:
                raise DocumentError('Word 压缩包包含异常路径。')
            seen.add(item.filename)
            if item.flag_bits & 1:
                raise DocumentError('Word 带有密码，请解除密码后重新上传。')
            if 'vbaproject' in item.filename.lower() or '/embeddings/' in item.filename.lower():
                raise DocumentError('Word 含宏或嵌入程序，请转成普通 DOCX 或 PDF。')
            if item.filename.endswith(('.xml', '.rels')):
                xml = archive.read(item)
                if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                    raise DocumentError('Word 含不支持的 XML 定义。')
                if item.filename.endswith('.rels'):
                    for relation in ET.fromstring(xml):
                        if relation.attrib.get('TargetMode', '').lower() == 'external':
                            raise DocumentError('Word 含外部链接资源，请移除链接或嵌入图片后上传。')
        if 'word/document.xml' not in seen:
            raise DocumentError('文件不是可读取的 Word 文档。')


def _docx(data, budget):
    _checked_docx(data, budget)
    doc = Document(BytesIO(data))
    result = ExtractedDocument(index_kind='paragraph')
    # 同一遍历涵盖普通段落和表格单元格，索引保留文档顺序，不虚构页码。
    for index, paragraph in enumerate(doc.element.body.iter(W + 'p'), 1):
        budget.check()
        if index > budget.limits.max_docx_paragraphs:
            raise DocumentError('Word 段落过多，请拆分文件。')
        text = budget.text(''.join(node.text or '' for node in paragraph.iter(W + 't')))
        formulas = []
        issues = []
        for formula in paragraph.iter(M + 'oMath'):
            # 保留OMML原结构用于后续公式转换；绝不伪称线性拼接等于原公式。
            formulas.append(budget.text(ET.tostring(formula, encoding='unicode')))
            issues.append('verify_formula')
        refs = []
        for blip in paragraph.iter(A + 'blip'):
            relation_id = blip.get(R + 'embed')
            if relation_id and relation_id in doc.part.rels:
                relation = doc.part.rels[relation_id]
                if relation.is_external:
                    raise DocumentError('Word 外部图片不能直接导入。')
                if relation.target_part.content_type not in IMAGE_FORMATS:
                    issues.append('unsupported_embedded_image')
                else:
                    # 不吞掉像素/资源限制错误，否则会把危险或不完整导入伪装为成功。
                    refs.append(_image_asset(relation.target_part.blob, result, budget, IMAGE_FORMATS[relation.target_part.content_type]))
            else:
                issues.append('missing_embedded_image')
        if list(paragraph.iter(W + 'pict')) or list(paragraph.iter(A + 'graphicData')) and not refs:
            issues.append('verify_word_drawing')
        result.sources.append(ExtractedSource(index, text, formulas, list(dict.fromkeys(refs)), list(dict.fromkeys(issues))))
    if any(str(part.partname).startswith(('/word/header', '/word/footer')) for part in doc.part.package.iter_parts()):
        result.warnings.append('headers_footers_not_extracted')
    return result


def extract_document(data: bytes, media_type: str, limits: DocumentLimits) -> ExtractedDocument:
    """只做本地解析；调用方须在有强制超时的工作进程中运行，不在HTTP线程直接解析。"""
    if media_type not in ('application/pdf', DOCX, *IMAGE_FORMATS):
        raise DocumentError('暂不支持此格式；旧版 DOC 请转为 DOCX 或 PDF。')
    maximum = limits.max_image_bytes if media_type in IMAGE_FORMATS else limits.max_document_bytes
    if not data or len(data) > maximum:
        raise DocumentError('文件为空或超过大小限制，请检查或分批上传。')
    budget = _Budget(limits)
    try:
        if media_type == 'application/pdf':
            return _pdf(data, budget)
        if media_type == DOCX:
            return _docx(data, budget)
        result = ExtractedDocument(index_kind='image')
        ref = _image_asset(data, result, budget, IMAGE_FORMATS[media_type])
        result.sources.append(ExtractedSource(1, '', asset_refs=[ref], warnings=['needs_ocr']))
        return result
    except DocumentError:
        raise
    except Exception:
        raise DocumentError('文件损坏或含暂不支持的内容，请重新导出文件后上传。') from None
