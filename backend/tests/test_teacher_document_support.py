"""本地导入能力使用合成教学文件验证，不外发真实材料。"""
import io
import unittest
import zipfile

from docx import Document
from docx.oxml import parse_xml
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from pypdf import PdfReader, PdfWriter

from app.teacher_assessment.extractors import DocumentLimits, DocumentError, extract_document


def sample_image():
    image = Image.new('RGB', (400, 240), 'white')
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((40, 40, 220, 180), outline='blue', width=4)
    drawing.text((45, 190), 'a = 5 cm', fill='black')
    result = io.BytesIO()
    image.save(result, format='PNG')
    return result.getvalue()


def sample_pdf():
    result = io.BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    doc = canvas.Canvas(result, pagesize=(300, 300))
    doc.setFont('STSong-Light', 14)
    doc.drawString(20, 260, '求正方形的面积，边长为5厘米。')
    doc.rect(50, 70, 130, 130)
    doc.showPage()
    from reportlab.lib.utils import ImageReader
    doc.drawImage(ImageReader(io.BytesIO(sample_image())), 20, 40, width=260, height=156)
    doc.save()
    return result.getvalue()


class TeacherDocumentSupportTests(unittest.TestCase):
    def test_mixed_pdf_preserves_chinese_text_and_scanned_page(self):
        extracted = extract_document(sample_pdf(), 'application/pdf', DocumentLimits())
        self.assertEqual([source.index for source in extracted.sources], [1, 2])
        self.assertIn('正方形', extracted.sources[0].text)
        self.assertEqual(extracted.sources[1].text.strip(), '')
        self.assertIn('needs_ocr', extracted.sources[1].warnings)
        self.assertIn('verify_visual_math', extracted.sources[0].warnings)
        for source in extracted.sources:
            self.assertTrue(source.asset_refs, '文字页和扫描页都保留可对照的原页')
            asset = extracted.assets[source.asset_refs[0]]
            with Image.open(io.BytesIO(asset.data)) as image:
                self.assertEqual(image.format, 'PNG')
                self.assertIsNotNone(image.getbbox())
                self.assertLess(image.convert('L').getextrema()[0], 200)

    def test_docx_preserves_table_formula_and_embedded_image_without_fake_pages(self):
        doc = Document()
        doc.add_paragraph('计算下式，并观察配图。')
        paragraph = doc.add_paragraph()
        paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'))
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '边长'
        table.cell(0, 1).text = '5厘米'
        doc.add_picture(io.BytesIO(sample_image()))
        data = io.BytesIO()
        doc.save(data)
        extracted = extract_document(data.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', DocumentLimits())
        combined = '\n'.join(source.text for source in extracted.sources)
        self.assertIn('边长', combined)
        self.assertIn('5厘米', combined)
        self.assertEqual(extracted.index_kind, 'paragraph')
        self.assertEqual([source.index for source in extracted.sources], list(range(1, len(extracted.sources)+1)))
        self.assertTrue(any(source.formulas for source in extracted.sources))
        self.assertTrue(any('verify_formula' in source.warnings for source in extracted.sources))
        self.assertTrue(any(source.asset_refs for source in extracted.sources))
        self.assertTrue(extracted.assets)

    def test_image_is_normalized_and_does_not_invent_ocr(self):
        extracted = extract_document(sample_image(), 'image/png', DocumentLimits())
        self.assertEqual(extracted.sources[0].text, '')
        self.assertIn('needs_ocr', extracted.sources[0].warnings)
        self.assertEqual(extracted.index_kind, 'image')

    def test_rotated_jpeg_respects_orientation_and_strips_metadata(self):
        image = Image.new('RGB', (120, 60), 'blue')
        metadata = Image.Exif()
        metadata[274] = 6
        metadata[270] = 'private camera note'
        data = io.BytesIO()
        image.save(data, format='JPEG', exif=metadata)
        extracted = extract_document(data.getvalue(), 'image/jpeg', DocumentLimits())
        asset = next(iter(extracted.assets.values()))
        with Image.open(io.BytesIO(asset.data)) as normalized:
            self.assertEqual(normalized.size, (60, 120))
            self.assertFalse(normalized.getexif())
            self.assertNotIn('private camera note', str(normalized.info))

    def test_docx_image_pixel_limit_is_not_swallowed_as_partial_success(self):
        doc = Document()
        doc.add_picture(io.BytesIO(sample_image()))
        data = io.BytesIO()
        doc.save(data)
        with self.assertRaisesRegex(DocumentError, '像素'):
            extract_document(data.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', DocumentLimits(max_image_pixels=100))

    def test_resource_limits_fail_explicitly_instead_of_truncating(self):
        with self.assertRaisesRegex(DocumentError, '大小'):
            extract_document(sample_image(), 'image/png', DocumentLimits(max_image_bytes=10))
        doc = Document()
        doc.add_paragraph('第一段')
        doc.add_paragraph('第二段')
        data = io.BytesIO()
        doc.save(data)
        for limits, message in [(DocumentLimits(max_docx_paragraphs=1), '段落'), (DocumentLimits(max_zip_entries=1), '解压'), (DocumentLimits(max_uncompressed_bytes=10), '解压')]:
            with self.subTest(message=message), self.assertRaisesRegex(DocumentError, message):
                extract_document(data.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', limits)

    def test_encrypted_pdf_and_page_limit_are_explicit_errors(self):
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(sample_pdf())))
        writer.encrypt('synthetic-password')
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(DocumentError, '密码'):
            extract_document(data.getvalue(), 'application/pdf', DocumentLimits())
        with self.assertRaisesRegex(DocumentError, '页数'):
            extract_document(sample_pdf(), 'application/pdf', DocumentLimits(max_pdf_pages=1))

    def test_unsupported_type_and_mismatched_signature_are_rejected(self):
        for data, media in [(b'<html>not an image</html>', 'image/png'), (sample_image(), 'application/pdf'), (b'legacy doc', 'application/msword')]:
            with self.subTest(media=media), self.assertRaises(DocumentError):
                extract_document(data, media, DocumentLimits())

    def test_docx_rejects_traversal_and_external_relationship(self):
        doc = Document()
        doc.add_paragraph('不应访问远程地址')
        data = io.BytesIO()
        doc.save(data)
        for entry, value in [('../escape.xml', b'x'), ('word/_rels/document.xml.rels', b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="https://example.invalid/private" TargetMode="External"/></Relationships>')]:
            target = io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(data.getvalue())) as original, zipfile.ZipFile(target, 'w') as changed:
                for item in original.infolist():
                    if item.filename != entry:
                        changed.writestr(item.filename, original.read(item))
                changed.writestr(entry, value)
            with self.subTest(entry=entry), self.assertRaises(DocumentError):
                extract_document(target.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', DocumentLimits())


if __name__ == '__main__':
    unittest.main()
