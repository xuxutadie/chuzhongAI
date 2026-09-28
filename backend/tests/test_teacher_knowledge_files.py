import unittest
from io import BytesIO
from pathlib import Path
from PIL import Image
from pypdf import PdfWriter
from docx import Document
from knowledge_test_support import KnowledgeFixture
from app.teacher_knowledge.files import KnowledgeFiles
from app.teacher_knowledge.repository import KnowledgeError


def png():
    buf = BytesIO()
    Image.new('RGB',(30,30),'white').save(buf,'PNG')
    return buf.getvalue()


class FileTests(KnowledgeFixture,unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.files = KnowledgeFiles(self.repo,Path(self.temp.name)/'private')

    def test_limits_and_private_reads(self):
        item = self.files.store(1,'题.png',png())
        duplicate = self.files.store(1,'题2.png',png())
        self.assertEqual(item['id'],duplicate['id'])
        self.assertNotEqual(item['id'],self.files.store(2,'题.png',png())['id'])
        with self.assertRaises(KnowledgeError) as caught:
            self.files.read(2,item['id'])
        self.assertEqual(caught.exception.status_code,404)
        for name,data in [('old.doc',b'a'),('fake.pdf',png()),('../a.png',png()),('big.png',b'x'*(5*1024*1024+1))]:
            with self.assertRaises(KnowledgeError):
                self.files.store(1,name,data)

    def test_source_positions_and_mixed_scan(self):
        book=Document()
        book.add_paragraph('第一节有理数')
        book.add_paragraph('两个数相加')
        buf=BytesIO(); book.save(buf)
        item=self.files.store(1,'教案.docx',buf.getvalue())
        result=self.files.extract(1,item['id'])
        self.assertEqual(result['index_kind'],'paragraph')
        self.assertIn('第一节有理数',result['sources'][0]['text'])
        self.assertEqual(result['chapter_suggestions'][0]['title'],'第一节有理数')
        image=self.files.store(1,'扫描.png',png())
        extracted=self.files.extract(1,image['id'])
        self.assertIn('needs_ocr',extracted['sources'][0]['warnings'])
        asset=extracted['sources'][0]['asset_refs'][0]
        self.assertTrue(self.repo.owned('files',1,asset)['data']['asset'])
        self.assertEqual(self.files.read(1,asset)[1],'image/png')
        with self.assertRaises(KnowledgeError): self.files.read(2,asset)

    def test_too_many_pages_and_encryption_rejected_on_extract(self):
        writer=PdfWriter()
        for _ in range(51): writer.add_blank_page(width=100,height=100)
        buf=BytesIO(); writer.write(buf)
        item=self.files.store(1,'长教材.pdf',buf.getvalue())
        with self.assertRaises(KnowledgeError): self.files.extract(1,item['id'])
        writer=PdfWriter(); writer.add_blank_page(width=100,height=100); writer.encrypt('secret')
        buf=BytesIO(); writer.write(buf)
        item=self.files.store(1,'加密.pdf',buf.getvalue())
        with self.assertRaises(KnowledgeError): self.files.extract(1,item['id'])

    def test_archive_and_restore(self):
        item=self.files.store(1,'图片.png',png())
        item=self.files.set_archived(1,item['id'],True,1)
        self.assertEqual(self.repo.listing('files',1)['total'],0)
        self.files.set_archived(1,item['id'],False,item['revision'])
        self.assertEqual(self.repo.listing('files',1)['total'],1)
