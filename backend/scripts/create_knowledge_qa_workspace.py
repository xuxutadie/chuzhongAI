"""新建隔离知识库验收环境，绝不覆盖已有库。"""
import argparse
import sqlite3
import sys
from pathlib import Path
from contextlib import closing
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.repositories.student_workspace_repository import StudentWorkspaceRepository
from app.repositories.teacher_schema import migrate_teacher_schema
from app.repositories.learning_route_repository import migrate_learning_schema
from app.teacher_knowledge.schema import migrate_knowledge_schema
from app.services.student_workspace_service import StudentWorkspaceService
from app.api.routes.student_workspace import get_student_workspace_service
from app.main import app
from fastapi.testclient import TestClient
from docx import Document


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--database',type=Path,required=True)
    target=parser.parse_args().database.resolve()
    if target.exists():parser.error('只允许新建验收数据库')
    target.parent.mkdir(parents=True,exist_ok=True)
    workspace=StudentWorkspaceService(StudentWorkspaceRepository(target))
    with closing(sqlite3.connect(target)) as db:
        db.execute('BEGIN IMMEDIATE');migrate_teacher_schema(db);migrate_learning_schema(db);migrate_knowledge_schema(db);db.commit()
    app.dependency_overrides[get_student_workspace_service]=lambda:workspace
    client=TestClient(app)
    for name in ['knowledge-qa-a','knowledge-qa-b']:
        response=client.post('/api/v1/auth/register-teacher',json={'username':name,'password':'Knowledge-QA-2026','display_name':'知识库验收教师（合成）','school_name':'验收学校','teaching_classes':['七年级1班']})
        assert response.status_code==201,response.status_code
    workspace.register_student(username='knowledge-qa-student',password='Knowledge-QA-2026',display_name='知识库验收学生（合成）')
    client.close();app.dependency_overrides.clear()
    doc=Document();doc.add_paragraph('有理数加法');doc.add_paragraph('有理数包括整数和分数。同号两数相加，取相同符号，并把绝对值相加。');doc.add_paragraph('计算 2 + 3，答案为 5。')
    doc.save(target.parent/'knowledge-qa.docx')
    print('隔离验收库与合成资料已建立，不包含真实学生信息。')
